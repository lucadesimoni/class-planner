"""Saubere Exporte: Klassenlisten und Stundenpläne als Excel und als druckfertiges HTML.

* klassenlisten.xlsx – Übersicht + ein Blatt pro Klasse (A4 hoch, Kopfzeile wiederholt)
* stundenplaene.xlsx – ein Blatt pro Klasse, Lehrperson und Raum (A4 quer) + alle Lektionen
* druck.html         – Klassenlisten und Stundenpläne, eine Seite pro Klasse/Lehrperson
                       (im Browser drucken oder als PDF speichern)

Förder- und Sonderschuldaten (IF, DaZ, IS) sind besonders schützenswerte Personendaten
und erscheinen nur mit ``mit_foerderdaten=True``.
"""
from __future__ import annotations

import html
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from .ablauf import Ergebnis

esc = html.escape


def neues_arbeitsbuch():
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _blattname(name: str, vergeben: set[str]) -> str:
    basis = re.sub(r"[\[\]:*?/\\]", "-", name)[:31] or "Blatt"
    kandidat, i = basis, 2
    while kandidat.lower() in vergeben:
        kandidat = f"{basis[:28]}~{i}"
        i += 1
    vergeben.add(kandidat.lower())
    return kandidat


def tabelle(ws, kopf: list[str], zeilen: list[list], start: int = 1) -> None:
    """Einfache, gut lesbare Tabelle mit fixierter Kopfzeile und Autofilter."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    duenn = Side(style="thin", color="C3C2B7")
    for c, wert in enumerate(kopf, 1):
        zelle = ws.cell(row=start, column=c, value=wert)
        zelle.font = Font(bold=True)
        zelle.fill = PatternFill("solid", fgColor="E1E0D9")
        zelle.border = Border(bottom=duenn)
    for r, zeile in enumerate(zeilen, start + 1):
        for c, wert in enumerate(zeile, 1):
            z = ws.cell(row=r, column=c, value=wert)
            z.alignment = Alignment(vertical="top", wrap_text=isinstance(wert, str) and "\n" in wert)
    for c, titel in enumerate(kopf, 1):
        breite = max([len(str(titel))] + [len(str(z[c - 1])) for z in zeilen if c - 1 < len(z)] + [4])
        ws.column_dimensions[ws.cell(row=start, column=c).column_letter].width = min(48, breite + 2)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    if zeilen:
        ws.auto_filter.ref = f"A{start}:{ws.cell(row=start + len(zeilen), column=len(kopf)).coordinate}"


def _druck(ws, quer: bool = False, kopfzeile: int | None = None) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if quer else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    if kopfzeile:
        ws.print_title_rows = f"{kopfzeile}:{kopfzeile}"
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.right.text = "Seite &P von &N"


def _namen(k) -> tuple[str, str]:
    if k.nachname or k.vorname:
        return k.nachname, k.vorname
    teile = k.name.rsplit(" ", 1)
    return (teile[1], teile[0]) if len(teile) == 2 else (k.name, "")


def _klassenlehrpersonen(r: Ergebnis) -> dict[str, list[str]]:
    lekt = defaultdict(lambda: defaultdict(int))
    for b in r.bloecke:
        if b.art == "unterricht":
            lekt[b.klasse][b.lehrperson] += b.lektionen
    return {k: [n for n, _ in sorted(v.items(), key=lambda t: -t[1])] for k, v in lekt.items()}


def _merkmale(k) -> str:
    return ", ".join(x for x, an in (("IF", k.if_), (f"DaZ {k.daz}".strip(), bool(k.daz)), ("IS", k.is_),
                                     ("Förderbedarf", k.foerderbedarf)) if an)


def klassenlisten_xlsx(r: Ergebnis, pfad: Path, mit_foerderdaten: bool = False) -> None:
    from openpyxl.styles import Font
    e, schule = r.einteilung, r.schule
    nach_id = {k.id: k for k in r.kinder}
    klp = _klassenlehrpersonen(r)
    wb = neues_arbeitsbuch()
    vergeben: set[str] = set()

    ws = wb.create_sheet(_blattname("Übersicht", vergeben))
    zeilen = []
    for c in e.klassen:
        ks = [nach_id[i] for i in c.kinder]
        zeilen.append([c.name, c.schulhaus, c.raum.name, "/".join(c.jahrgaenge), len(ks),
                       sum(k.geschlecht == "w" for k in ks), sum(k.geschlecht == "m" for k in ks),
                       ", ".join(klp.get(c.name, [])[:2])])
    tabelle(ws, ["Klasse", "Schulhaus", "Schulzimmer", "Jahrgänge", "Lernende", "Mädchen", "Knaben", "Klassenlehrperson(en)"],
            zeilen)
    _druck(ws, kopfzeile=1)

    kopf = ["Nr.", "Nachname", "Vorname", "Jahrgang", "Geschlecht", "Adresse"] + (
        ["Förderung"] if mit_foerderdaten else [])
    for c in e.klassen:
        ws = wb.create_sheet(_blattname(c.name, vergeben))
        ws["A1"] = f"Klassenliste {c.name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (f"{schule.name} · {schule.schuljahr + ' · ' if schule.schuljahr else ''}Schulhaus {c.schulhaus}, "
                    f"Schulzimmer {c.raum.name}")
        ws["A3"] = f"Klassenlehrperson: {', '.join(klp.get(c.name, [])[:2]) or '–'} · {len(c.kinder)} Lernende"
        ks = sorted((nach_id[i] for i in c.kinder), key=lambda k: (_namen(k)[0], _namen(k)[1]))
        zeilen = [[i, *_namen(k), k.jahrgang, k.geschlecht, k.adresse] + ([_merkmale(k)] if mit_foerderdaten else [])
                  for i, k in enumerate(ks, 1)]
        tabelle(ws, kopf, zeilen, start=5)
        _druck(ws, kopfzeile=5)
        ws.oddHeader.right.text = f"Stand {date.today():%d.%m.%Y}"
    wb.save(pfad)


def _raster(schule):
    tage = list(dict.fromkeys(s.tag for s in schule.slots))
    halbtage = list(dict.fromkeys(s.halbtag for s in schule.slots))
    zeilen = sorted(dict.fromkeys((s.halbtag, s.nr, s.zeit) for s in schule.slots),
                    key=lambda z: (halbtage.index(z[0]), z[1]))
    idx = {(s.tag, s.halbtag, s.nr): s.index for s in schule.slots}
    return tage, zeilen, idx


def stundenplaene_xlsx(r: Ergebnis, pfad: Path) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    schule = r.schule
    tage, zeilen, idx = _raster(schule)
    duenn = Side(style="thin", color="C3C2B7")
    wb = neues_arbeitsbuch()
    vergeben: set[str] = set()

    def plan(titel: str, untertitel: str, eintraege: dict[int, list[str]], gesperrt: set[int]):
        ws = wb.create_sheet(_blattname(titel, vergeben))
        ws["A1"] = titel
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = untertitel
        for c, tag in enumerate(tage, 2):
            z = ws.cell(row=4, column=c, value=tag)
            z.font = Font(bold=True)
            z.fill = PatternFill("solid", fgColor="E1E0D9")
            z.alignment = Alignment(horizontal="center")
            ws.column_dimensions[z.column_letter].width = 24
        ws.column_dimensions["A"].width = 8
        for r_, (ht, nr, zeit) in enumerate(zeilen, 5):
            ws.cell(row=r_, column=1, value=zeit).font = Font(bold=True)
            ws.row_dimensions[r_].height = 48
            for c, tag in enumerate(tage, 2):
                i = idx.get((tag, ht, nr))
                z = ws.cell(row=r_, column=c, value="\n".join(eintraege.get(i, [])) if i is not None else "")
                z.alignment = Alignment(wrap_text=True, vertical="top")
                z.border = Border(top=duenn, bottom=duenn, left=duenn, right=duenn)
                if i is None or (i in gesperrt and i not in eintraege):
                    z.fill = PatternFill("solid", fgColor="F0EFEC")
        _druck(ws, quer=True)

    for c in r.einteilung.klassen:
        eintr = defaultdict(list)
        for l in sorted(r.lektionen, key=lambda l: l.art != "unterricht"):
            if c.name in l.klassen:
                eintr[l.slot].append(f"{l.fach} – {l.lehrperson} ({l.raum_label})" if l.art == "unterricht"
                                     else f"+ {l.fach}: {l.lehrperson}")
        plan(f"Klasse {c.name}", f"Schulhaus {c.schulhaus}, Schulzimmer {c.raum.name}", eintr,
             schule.schulhaus(c.schulhaus).nicht_verfuegbar)
    for lp in schule.lehrpersonen:
        eintr = defaultdict(list)
        for l in r.lektionen:
            if l.lehrperson == lp.name:
                eintr[l.slot].append(f"{'/'.join(l.klassen)}: {l.fach}" + (f" ({l.raum_label})" if l.raum else ""))
        if eintr:
            plan(lp.name, f"{sum(len(v) for v in eintr.values())} Lektionen", eintr, lp.nicht_verfuegbar)
    # Spezialräume (alle Räume, die nicht Stammzimmer einer Klasse sind)
    stammzimmer = {(c.schulhaus, c.raum.name) for c in r.einteilung.klassen}
    raeume = defaultdict(lambda: defaultdict(list))
    for l in r.lektionen:
        if l.raum and (l.raum_ort, l.raum) not in stammzimmer:
            raeume[l.raum_ort, l.raum][l.slot].append(f"{'/'.join(l.klassen)}: {l.fach}")
    for (ort, raum), eintr in sorted(raeume.items()):
        plan(f"Raum {raum}", f"Schulhaus {ort}", eintr, set())
    ws = wb.create_sheet(_blattname("Alle Lektionen", vergeben))
    tabelle(ws, ["Tag", "Zeit", "Klasse", "Fach", "Art", "Lehrperson", "Raum", "Schulhaus"],
            [[schule.slots[l.slot].tag, schule.slots[l.slot].zeit, l.klasse, l.fach, l.art, l.lehrperson,
              l.raum_label, l.schulhaus] for l in sorted(r.lektionen, key=lambda l: (l.slot, l.klasse))])
    wb.save(pfad)


DRUCK_CSS = """
@page { size: A4; margin: 14mm 12mm; }
body { font: 11pt/1.35 system-ui,-apple-system,"Segoe UI",sans-serif; color:#0b0b0b; margin:0 }
.seite { page-break-after: always; padding: 0 0 8mm }
.seite.quer { page: quer } @page quer { size: A4 landscape; }
h1 { font-size: 16pt; margin: 0 0 2mm } .sub { color:#52514e; margin:0 0 5mm }
table { border-collapse: collapse; width: 100% } th, td { border: 1px solid #c3c2b7; padding: 3px 5px; text-align:left;
  vertical-align: top } th { background:#f0efec }
.plan td { height: 15mm; width: 18%; font-size: 9pt } .plan td.zu { background:#f0efec }
.plan .foerder { color:#52514e; font-style: italic }
.fuss { color:#898781; font-size: 8pt; margin-top: 4mm }
@media screen { body { max-width: 1000px; margin: 0 auto; padding: 16px; background:#f9f9f7 }
  .seite { background:#fff; border:1px solid #e1e0d9; margin: 0 0 16px; padding: 16px } }
"""


def druck_html(r: Ergebnis, pfad: Path, mit_foerderdaten: bool = False) -> None:
    schule, e = r.schule, r.einteilung
    nach_id = {k.id: k for k in r.kinder}
    klp = _klassenlehrpersonen(r)
    tage, zeilen, idx = _raster(schule)
    stand = f"Stand {date.today():%d.%m.%Y}"

    def plan(eintraege: dict[int, str], gesperrt: set[int]) -> str:
        out = ["<table class='plan'><tr><th></th>" + "".join(f"<th>{t}</th>" for t in tage) + "</tr>"]
        for ht, nr, zeit in zeilen:
            out.append(f"<tr><th>{zeit}</th>")
            for t in tage:
                i = idx.get((t, ht, nr))
                zu = " class='zu'" if i is None or (i in gesperrt and i not in eintraege) else ""
                out.append(f"<td{zu}>{eintraege.get(i, '')}</td>")
            out.append("</tr>")
        return "".join(out) + "</table>"

    teile = [f"<!doctype html><html lang='de'><head><meta charset='utf-8'><title>{esc(schule.name)} – Listen und "
             f"Stundenpläne</title><style>{DRUCK_CSS}</style></head><body>"]
    for c in e.klassen:
        ks = sorted((nach_id[i] for i in c.kinder), key=lambda k: _namen(k))
        kopf = "<tr><th>Nr.</th><th>Nachname</th><th>Vorname</th><th>Jg.</th><th>G.</th><th>Adresse</th>" + (
            "<th>Förderung</th>" if mit_foerderdaten else "") + "</tr>"
        zeilen_html = "".join(
            f"<tr><td>{i}</td><td>{esc(_namen(k)[0])}</td><td>{esc(_namen(k)[1])}</td><td>{esc(k.jahrgang)}</td>"
            f"<td>{esc(k.geschlecht)}</td><td>{esc(k.adresse)}</td>"
            + (f"<td>{esc(_merkmale(k))}</td>" if mit_foerderdaten else "") + "</tr>" for i, k in enumerate(ks, 1))
        teile.append(f"<section class='seite'><h1>Klassenliste {esc(c.name)}</h1><p class='sub'>{esc(schule.name)} · "
                     f"Schulhaus {esc(c.schulhaus)}, Zimmer {esc(c.raum.name)} · Klassenlehrperson: "
                     f"{esc(', '.join(klp.get(c.name, [])[:2]) or '–')} · {len(ks)} Lernende</p>"
                     f"<table>{kopf}{zeilen_html}</table><p class='fuss'>{stand}</p></section>")
        eintr = defaultdict(str)
        for l in sorted(r.lektionen, key=lambda l: l.art != "unterricht"):
            if c.name in l.klassen:
                eintr[l.slot] += (f"<b>{esc(l.fach)}</b><br>{esc(l.lehrperson)}<br>{esc(l.raum_label)}"
                                  if l.art == "unterricht" else
                                  f"<div class='foerder'>+ {esc(l.fach)}: {esc(l.lehrperson)}</div>")
        teile.append(f"<section class='seite quer'><h1>Stundenplan {esc(c.name)}</h1><p class='sub'>"
                     f"{esc(schule.name)} · Schulhaus {esc(c.schulhaus)}</p>"
                     f"{plan(eintr, schule.schulhaus(c.schulhaus).nicht_verfuegbar)}<p class='fuss'>{stand}</p></section>")
    for lp in schule.lehrpersonen:
        eintr = {l.slot: f"<b>{esc('/'.join(l.klassen))}</b><br>{esc(l.fach)}<br>{esc(l.raum_label)}"
                 for l in r.lektionen if l.lehrperson == lp.name}
        if eintr:
            teile.append(f"<section class='seite quer'><h1>Stundenplan {esc(lp.name)}</h1><p class='sub'>"
                         f"{len(eintr)} Lektionen</p>{plan(eintr, lp.nicht_verfuegbar)}<p class='fuss'>{stand}</p>"
                         "</section>")
    teile.append("</body></html>")
    pfad.write_text("".join(teile), encoding="utf-8")


def exportiere_listen(r: Ergebnis, ordner: Path, mit_foerderdaten: bool = False) -> list[Path]:
    if not r.einteilung:
        return []
    dateien = [ordner / "klassenlisten.xlsx"]
    klassenlisten_xlsx(r, dateien[0], mit_foerderdaten)
    if r.lektionen:
        dateien += [ordner / "stundenplaene.xlsx", ordner / "druck.html"]
        stundenplaene_xlsx(r, dateien[1])
        druck_html(r, dateien[2], mit_foerderdaten)
    return dateien
