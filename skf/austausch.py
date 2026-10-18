# SPDX-License-Identifier: LicenseRef-Proprietary
# Copyright (c) 2026 Outthinkx for Communities. Alle Rechte vorbehalten.
"""Standardisierter Import/Export von Konfiguration, Stammdaten und Kinderlisten.

Formate:
  * Kinderliste: CSV (Komma oder Semikolon, UTF-8 oder Windows-1252) oder Excel (.xlsx).
    Spaltennamen werden tolerant erkannt (Vorname/Nachname, Strasse/PLZ/Ort,
    WGS84 lat/lon oder Schweizer LV95-Koordinaten E/N, Ja/Nein-Werte, ...).
  * Stammdaten (Schule + Kinder) als ein Excel-Arbeitsbuch mit einem Blatt pro Tabelle
    oder als JSON-Datei (Format "skf-stammdaten", Schema in schema/).
Import und Export sind verlustfrei: Export → Import ergibt dieselbe Konfiguration.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date
from pathlib import Path

import yaml

from .daten import Kind

FORMAT = "skf-stammdaten"
VERSION = 1
JA = {"1", "ja", "j", "yes", "y", "true", "wahr", "x", "✓"}

# Standardspalten der Kinderliste (Export) – Reihenfolge = Spaltenreihenfolge
KINDER_SPALTEN = ["id", "vorname", "nachname", "name", "jahrgang", "geschlecht", "adresse", "lat", "lon",
                  "familie", "if", "daz", "is", "foerderbedarf", "schulhaus_fix", "bisheriges_schulhaus",
                  "bisherige_klasse", "zusammen_mit", "nicht_mit", "wunsch_schulhaus", "bemerkung"]

# erkannte Spaltennamen (normalisiert: klein, Umlaute aufgelöst, nur a-z0-9)
ALIASE = {
    "id": ["id", "schuelerid", "sid", "lernendenid", "personid", "nummer", "schuelernummer", "nr"],
    "vorname": ["vorname", "firstname", "rufname"],
    "nachname": ["nachname", "familienname", "lastname", "surname"],
    "name": ["name", "vollername", "fullname"],
    "jahrgang": ["jahrgang", "stufe", "klassenstufe", "schuljahrneu", "jg", "jahrgangneu", "stufeneu"],
    "geschlecht": ["geschlecht", "sex", "gender", "mw"],
    "adresse": ["adresse", "address", "wohnadresse"],
    "strasse": ["strasse", "str", "street", "strassenr", "strasseundnr"],
    "hausnummer": ["hausnummer", "hausnr"],
    "plz": ["plz", "postleitzahl", "zip"],
    "ort": ["ort", "wohnort", "city", "gemeinde"],
    "lat": ["lat", "breite", "latitude", "breitengrad"],
    "lon": ["lon", "lng", "laenge", "longitude", "laengengrad"],
    "e": ["e", "lv95e", "ekoordinate", "koordinatee", "ost", "gkodex", "x"],
    "n": ["n", "lv95n", "nkoordinate", "koordinaten", "nord", "gkodey", "y"],
    "familie": ["familie", "familienid", "haushalt", "haushaltid", "geschwistergruppe"],
    "if": ["if", "integrativefoerderung"],
    "daz": ["daz", "deutschalszweitsprache"],
    "is": ["is", "integrativesonderschulung", "sonderschulung"],
    "foerderbedarf": ["foerderbedarf", "foerderung", "besondererbedarf"],
    "schulhaus_fix": ["schulhausfix", "fixesschulhaus", "zuteilungfix"],
    "bisheriges_schulhaus": ["bisherigesschulhaus", "schulhausvorjahr", "schulhausbisher"],
    "bisherige_klasse": ["bisherigeklasse", "klassevorjahr", "klassebisher", "kggruppe", "kindergartengruppe",
                         "kigagruppe", "kindergartenvorjahr"],
    "zusammen_mit": ["zusammenmit", "wunschzusammen", "wunschmit", "gspaenli", "freundschaftswunsch",
                     "wunschmitschueler"],
    "nicht_mit": ["nichtmit", "nichtzusammenmit", "trennenvon", "wunschgetrennt"],
    "wunsch_schulhaus": ["wunschschulhaus", "schulhauswunsch"],
    "bemerkung": ["bemerkung", "bemerkungen", "notiz", "kommentar"],
}


def _norm(text: str) -> str:
    t = str(text).strip().lower()
    for a, b in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss"), ("é", "e"), ("è", "e")):
        t = t.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", t)


def lv95_zu_wgs84(e: float, n: float) -> tuple[float, float]:
    """Schweizer Landeskoordinaten LV95 (E/N) → WGS84 (lat, lon).

    Näherungsformel von swisstopo, Genauigkeit ca. 1 m – für Schulwege mehr als genug.
    """
    y, x = (e - 2_600_000) / 1e6, (n - 1_200_000) / 1e6
    lam = 2.6779094 + 4.728982 * y + 0.791484 * y * x + 0.1306 * y * x * x - 0.0436 * y ** 3
    phi = 16.9023892 + 3.238272 * x - 0.270978 * y * y - 0.002528 * x * x - 0.0447 * y * y * x - 0.0140 * x ** 3
    return phi * 100 / 36, lam * 100 / 36


def normalisiere_jahrgang(wert: str) -> str:
    """"1. Klasse" → "1", "KG 2" / "2. Kindergarten" → "KG2"; Unbekanntes bleibt unverändert."""
    t = str(wert).strip()
    klein = t.lower().replace(" ", "")
    ziffer = re.search(r"\d+", klein)
    if ("kg" in klein or "kindergarten" in klein) and ziffer:
        return f"KG{ziffer.group()}"
    if re.fullmatch(r"\d+\.?(klasse|kl\.?|primar|p)?", klein):
        return ziffer.group()
    return t


def normalisiere_geschlecht(wert: str) -> str:
    t = _norm(wert)
    if t in ("w", "f", "weiblich", "female", "maedchen", "frau"):
        return "w"
    if t in ("m", "maennlich", "male", "knabe", "junge", "mann"):
        return "m"
    return "d" if t in ("d", "divers", "x") else ""


def _ja(wert) -> bool:
    t = str(wert or "").strip().lower()
    return t in JA or (_norm(t) in {_norm(j) for j in JA} - {""})


def _zeilen(pfad: str | Path, blaetter: tuple[str, ...] = ("Kinder",)) -> list[dict[str, str]]:
    pfad = Path(pfad)
    if pfad.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(pfad, read_only=True, data_only=True)
        passend = [n for n in wb.sheetnames if _norm(n) in {_norm(b) for b in blaetter}]
        ws = wb[passend[0]] if passend else wb.worksheets[0]
        zeilen = list(ws.iter_rows(values_only=True))
        kopf = [str(k or "").strip() for k in zeilen[0]]
        return [{kopf[i]: ("" if v is None else str(v)) for i, v in enumerate(z) if i < len(kopf)}
                for z in zeilen[1:] if any(v not in (None, "") for v in z)]
    roh = pfad.read_bytes()
    for kodierung in ("utf-8-sig", "cp1252"):
        try:
            text = roh.decode(kodierung)
            break
        except UnicodeDecodeError:
            continue
    erste = text.splitlines()[0] if text else ""
    trenner = ";" if erste.count(";") > erste.count(",") else ("\t" if "\t" in erste else ",")
    return list(csv.DictReader(io.StringIO(text), delimiter=trenner))


def lese_kinder(pfad: str | Path, spalten: dict[str, str] | None = None) -> list[Kind]:
    """Kinderliste laden. ``spalten`` ordnet bei Bedarf eigene Spaltennamen zu,
    z.B. {"id": "Schüler-Nr", "jahrgang": "Stufe 26/27"}."""
    zeilen = _zeilen(pfad)
    if not zeilen:
        return []
    kopf = list(zeilen[0].keys())
    zuordnung: dict[str, str] = {}
    for feld, quelle in (spalten or {}).items():
        if quelle not in kopf:
            raise ValueError(f"Spalte '{quelle}' (für {feld}) fehlt in {pfad}")
        zuordnung[feld] = quelle
    normiert = {_norm(k): k for k in kopf}
    for feld, aliase in ALIASE.items():
        if feld not in zuordnung:
            treffer = next((normiert[a] for a in aliase if a in normiert and normiert[a] not in zuordnung.values()),
                           None)
            if treffer:
                zuordnung[feld] = treffer
    if "id" not in zuordnung or "jahrgang" not in zuordnung:
        raise ValueError(f"{pfad}: Spalten für 'id' und 'jahrgang' nicht gefunden (vorhanden: {', '.join(kopf)}). "
                         "Mit --spalte id=<Spaltenname> zuordnen.")
    benutzt = set(zuordnung.values())

    kinder = []
    for zeile in zeilen:
        def w(feld: str) -> str:
            return str(zeile.get(zuordnung.get(feld, ""), "") or "").strip()

        def zahl(feld: str) -> float | None:
            v = w(feld).replace("'", "").replace(" ", "")
            return float(v.replace(",", ".")) if v else None

        lat, lon = zahl("lat"), zahl("lon")
        if lat is None and zahl("e") is not None and zahl("n") is not None:
            e, n = zahl("e"), zahl("n")
            if e < 2_000_000:  # LV03 → LV95
                e, n = e + 2_000_000, n + 1_000_000
            lat, lon = lv95_zu_wgs84(e, n)
        adresse = w("adresse")
        if not adresse and (w("strasse") or w("ort")):
            strasse = " ".join(x for x in (w("strasse"), w("hausnummer")) if x)
            adresse = ", ".join(x for x in (strasse, " ".join(x for x in (w("plz"), w("ort")) if x)) if x)
        name = w("name") or " ".join(x for x in (w("vorname"), w("nachname")) if x)
        daz = _norm(w("daz"))
        daz = "aufbau" if _ja(w("daz")) else ("" if daz in ("", "0", "nein", "n", "no", "false") else daz)
        kinder.append(Kind(
            id=w("id"), name=name, jahrgang=normalisiere_jahrgang(w("jahrgang")),
            geschlecht=normalisiere_geschlecht(w("geschlecht")), adresse=adresse, lat=lat, lon=lon,
            familie=w("familie"), foerderbedarf=_ja(w("foerderbedarf")), schulhaus_fix=w("schulhaus_fix"),
            nicht_mit=[x.strip() for x in re.split(r"[;,]", w("nicht_mit")) if x.strip()],
            if_=_ja(w("if")), daz=daz, is_=_ja(w("is")),
            bisheriges_schulhaus=w("bisheriges_schulhaus"), bisherige_klasse=w("bisherige_klasse"),
            zusammen_mit=[x.strip() for x in re.split(r"[;,]", w("zusammen_mit")) if x.strip()],
            wunsch_schulhaus=w("wunsch_schulhaus"), bemerkung=w("bemerkung"),
            vorname=w("vorname"), nachname=w("nachname"),
            zusatz={k: str(v) for k, v in zeile.items() if k not in benutzt and k and v not in (None, "")}))
    return kinder


def kind_als_dict(k: Kind) -> dict[str, str]:
    d = {
        "id": k.id, "vorname": k.vorname, "nachname": k.nachname, "name": k.name, "jahrgang": k.jahrgang,
        "geschlecht": k.geschlecht, "adresse": k.adresse,
        "lat": "" if k.lat is None else f"{k.lat:.6f}", "lon": "" if k.lon is None else f"{k.lon:.6f}",
        "familie": k.familie, "if": "ja" if k.if_ else "", "daz": k.daz, "is": "ja" if k.is_ else "",
        "foerderbedarf": "ja" if k.foerderbedarf else "", "schulhaus_fix": k.schulhaus_fix,
        "bisheriges_schulhaus": k.bisheriges_schulhaus, "bisherige_klasse": k.bisherige_klasse,
        "zusammen_mit": ";".join(k.zusammen_mit), "nicht_mit": ";".join(k.nicht_mit),
        "wunsch_schulhaus": k.wunsch_schulhaus, "bemerkung": k.bemerkung,
    }
    d.update({f: v for f, v in k.zusatz.items() if f not in d})
    return d


def schreibe_kinder(kinder: list[Kind], pfad: str | Path) -> None:
    pfad = Path(pfad)
    zeilen = [kind_als_dict(k) for k in kinder]
    spalten = KINDER_SPALTEN + sorted({f for z in zeilen for f in z} - set(KINDER_SPALTEN))
    if pfad.suffix.lower() == ".xlsx":
        from .export import neues_arbeitsbuch, tabelle
        wb = neues_arbeitsbuch()
        tabelle(wb.create_sheet("Kinder"), spalten, [[z.get(s, "") for s in spalten] for z in zeilen])
        wb.save(pfad)
        return
    with open(pfad, "w", newline="", encoding="utf-8") as f:
        wr = csv.DictWriter(f, fieldnames=spalten)
        wr.writeheader()
        wr.writerows(zeilen)


# ---------------------------------------------------------------- Stammdaten

def _liste(v) -> str:
    if v is None:
        return ""
    return ";".join(str(x) for x in v) if isinstance(v, list) else str(v)


def _aus_liste(v) -> list[str] | None:
    if v in (None, ""):
        return None
    return [x.strip() for x in str(v).split(";") if x.strip()]


def _zahl(v):
    if v in (None, ""):
        return None
    f = float(v)
    return int(f) if f.is_integer() else f


BLAETTER = {
    "Schulhaeuser": ["name", "kuerzel", "lat", "lon", "nicht_verfuegbar", "nutzt"],
    "Raeume": ["schulhaus", "name", "typ", "anzahl", "plaetze", "nicht_verfuegbar"],
    "Klassentypen": ["jahrgaenge", "bezeichnung", "raumtyp", "min", "max", "ideal", "strafe", "min_pro_jahrgang",
                     "schulhaeuser"],
    "Faecher": ["name", "raumtyp", "dauer", "max_pro_tag"],
    "Lehrpersonen": ["name", "kuerzel", "faecher", "jahrgaenge", "schulhaeuser", "pensum_min", "pensum_max",
                     "nicht_verfuegbar", "bisherige_klasse"],
    "Foerderangebote": ["name", "merkmal", "basis", "pro_kind", "frei_bis", "max", "parallel_zu", "raumtyp",
                        "quelle"],
}


def lese_rohkonfiguration(pfad: str | Path) -> dict:
    return yaml.safe_load(Path(pfad).read_text(encoding="utf-8"))


def schreibe_rohkonfiguration(cfg: dict, pfad: str | Path) -> None:
    Path(pfad).write_text(yaml.safe_dump(cfg, allow_unicode=True, sort_keys=False, width=120), encoding="utf-8")


def _tabellen_aus_cfg(cfg: dict) -> tuple[dict[str, list[dict]], dict]:
    """Zerlegt die Konfiguration in Tabellen + Resteinstellungen."""
    rest = json.loads(json.dumps(cfg))  # tiefe Kopie
    t: dict[str, list[dict]] = {k: [] for k in BLAETTER}
    for sh in rest.pop("schulhaeuser", []):
        for r in sh.pop("raeume", []):
            t["Raeume"].append({"schulhaus": sh["name"], **r})
        sh["nutzt"] = ";".join(f"{k}={v}" for k, v in (sh.get("nutzt") or {}).items())
        t["Schulhaeuser"].append(sh)
    kb = rest.setdefault("klassenbildung", {})
    t["Klassentypen"] = kb.pop("vorlagen", [])
    t["Faecher"] = [{"name": n, **(f or {})} for n, f in rest.pop("faecher", {}).items()]
    for lp in rest.pop("lehrpersonen", []):
        pensum = lp.pop("pensum", None)
        if pensum is not None:
            pensum = [pensum, pensum] if isinstance(pensum, int) else pensum
            lp["pensum_min"], lp["pensum_max"] = pensum
        t["Lehrpersonen"].append(lp)
    sp = rest.get("sonderpaedagogik") or {}
    for a in sp.pop("foerderangebote", []) or []:
        b = a.get("basis")
        if isinstance(b, dict):
            a["basis"] = ";".join(f"{k}={v}" for k, v in b.items())
        t["Foerderangebote"].append(a)
    stundentafel = rest.pop("stundentafel", {})
    return t, {"rest": rest, "stundentafel": stundentafel}


def _cfg_aus_tabellen(t: dict[str, list[dict]], rest: dict, stundentafel: dict) -> dict:
    cfg = json.loads(json.dumps(rest))
    raeume: dict[str, list] = {}
    for r in t.get("Raeume", []):
        sh = r.pop("schulhaus")
        raeume.setdefault(sh, []).append({k: v for k, v in r.items() if v not in (None, "")})
    schulhaeuser = []
    for sh in t.get("Schulhaeuser", []):
        e = {k: v for k, v in sh.items() if v not in (None, "")}
        if e.get("nutzt"):
            e["nutzt"] = dict(x.split("=", 1) for x in str(e["nutzt"]).split(";") if "=" in x)
        else:
            e.pop("nutzt", None)
        e["raeume"] = raeume.get(e["name"], [])
        schulhaeuser.append(e)
    cfg["schulhaeuser"] = schulhaeuser
    cfg.setdefault("klassenbildung", {})["vorlagen"] = [
        {k: v for k, v in v.items() if v not in (None, "")} for v in t.get("Klassentypen", [])]
    cfg["faecher"] = {f.pop("name"): {k: v for k, v in f.items() if v not in (None, "")} or None
                      for f in t.get("Faecher", [])}
    cfg["stundentafel"] = stundentafel
    lehrpersonen = []
    for lp in t.get("Lehrpersonen", []):
        e = {k: v for k, v in lp.items() if v not in (None, "")}
        if "pensum_min" in e or "pensum_max" in e:
            e["pensum"] = [e.pop("pensum_min", 0), e.pop("pensum_max", 99)]
        lehrpersonen.append(e)
    cfg["lehrpersonen"] = lehrpersonen
    angebote = []
    for a in t.get("Foerderangebote", []):
        e = {k: v for k, v in a.items() if v not in (None, "")}
        if isinstance(e.get("basis"), str) and "=" in e["basis"]:
            e["basis"] = {k: _zahl(v) for k, v in (x.split("=", 1) for x in e["basis"].split(";"))}
        angebote.append(e)
    if angebote:
        cfg.setdefault("sonderpaedagogik", {})["foerderangebote"] = angebote
    return cfg


LISTENFELDER = {"nicht_verfuegbar", "jahrgaenge", "schulhaeuser", "faecher", "parallel_zu"}
ZAHLFELDER = {"lat", "lon", "anzahl", "plaetze", "min", "max", "ideal", "strafe", "min_pro_jahrgang", "dauer",
              "max_pro_tag", "pensum_min", "pensum_max", "pro_kind", "frei_bis"}


def exportiere_stammdaten(cfg: dict, kinder: list[Kind], pfad: str | Path) -> None:
    """Schule + Kinder als Excel-Arbeitsbuch (.xlsx) oder JSON (.json) speichern."""
    pfad = Path(pfad)
    if pfad.suffix.lower() == ".json":
        pfad.write_text(json.dumps({"format": FORMAT, "version": VERSION, "exportiert": date.today().isoformat(),
                                    "schule": cfg, "kinder": [kind_als_dict(k) for k in kinder]},
                                   ensure_ascii=False, indent=1), encoding="utf-8")
        return
    from .export import neues_arbeitsbuch, tabelle
    t, extra = _tabellen_aus_cfg(cfg)
    wb = neues_arbeitsbuch()
    info = wb.create_sheet("Info")
    tabelle(info, ["Feld", "Wert"], [["Format", FORMAT], ["Version", VERSION], ["Schule", cfg.get("name", "")],
                                      ["Exportiert", date.today().isoformat()],
                                      ["Hinweis", "Listen mit ; trennen. Verfügbarkeit z.B. Mi;Fr-NM;Mo-VM-1. "
                                                  "Blatt 'Einstellungen' enthält die übrigen Angaben als YAML."]])
    for blatt, spalten in BLAETTER.items():
        extra_spalten = sorted({k for z in t[blatt] for k in z} - set(spalten))
        alle = spalten + extra_spalten
        tabelle(wb.create_sheet(blatt), alle, [[_liste(z.get(s)) if s in LISTENFELDER or isinstance(z.get(s), list)
                                                 else ("" if z.get(s) is None else z.get(s)) for s in alle]
                                                for z in t[blatt]])
    tafel = extra["stundentafel"]
    stufen = list(tafel)
    faecher = list(dict.fromkeys(f for st in tafel.values() for f in st))
    tabelle(wb.create_sheet("Stundentafel"), ["Fach"] + stufen,
            [[f] + [tafel[s].get(f, "") for s in stufen] for f in faecher])
    zeilen = [kind_als_dict(k) for k in kinder]
    spalten = KINDER_SPALTEN + sorted({f for z in zeilen for f in z} - set(KINDER_SPALTEN))
    tabelle(wb.create_sheet("Kinder"), spalten, [[z.get(s, "") for s in spalten] for z in zeilen])
    tabelle(wb.create_sheet("Einstellungen"), ["Schluessel", "YAML"],
            [[k, yaml.safe_dump(v, allow_unicode=True, sort_keys=False).strip()] for k, v in extra["rest"].items()])
    wb.save(pfad)


def importiere_stammdaten(pfad: str | Path) -> tuple[dict, list[Kind]]:
    """Gegenstück zu exportiere_stammdaten: liefert (Konfiguration, Kinder)."""
    pfad = Path(pfad)
    if pfad.suffix.lower() == ".json":
        daten = json.loads(pfad.read_text(encoding="utf-8"))
        if daten.get("format") != FORMAT:
            raise ValueError(f"{pfad}: kein {FORMAT}-Format")
        tmp = pfad.with_suffix(".kinder.tmp.csv")
        try:
            with open(tmp, "w", newline="", encoding="utf-8") as f:
                spalten = list(dict.fromkeys(k for z in daten["kinder"] for k in z)) or KINDER_SPALTEN
                wr = csv.DictWriter(f, fieldnames=spalten)
                wr.writeheader()
                wr.writerows(daten["kinder"])
            kinder = lese_kinder(tmp)
        finally:
            tmp.unlink(missing_ok=True)
        return daten["schule"], kinder

    from openpyxl import load_workbook
    wb = load_workbook(pfad, data_only=True)

    def blatt(name: str) -> list[dict]:
        if name not in wb.sheetnames:
            return []
        zeilen = list(wb[name].iter_rows(values_only=True))
        if not zeilen:
            return []
        kopf = [str(k) for k in zeilen[0]]
        out = []
        for z in zeilen[1:]:
            if all(v in (None, "") for v in z):
                continue
            e = {}
            for k, v in zip(kopf, z):
                if v in (None, ""):
                    continue
                if k in LISTENFELDER:
                    v = _aus_liste(v)
                elif k in ZAHLFELDER:
                    v = _zahl(v)
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                e[k] = v
            out.append(e)
        return out

    t = {name: blatt(name) for name in BLAETTER}
    stundentafel: dict[str, dict] = {}
    if "Stundentafel" in wb.sheetnames:
        zeilen = list(wb["Stundentafel"].iter_rows(values_only=True))
        stufen = [str(s) for s in zeilen[0][1:]]
        for z in zeilen[1:]:
            for s, v in zip(stufen, z[1:]):
                if v not in (None, ""):
                    stundentafel.setdefault(s, {})[str(z[0])] = int(v)
    rest = {str(z[0]): yaml.safe_load(str(z[1])) for z in list(wb["Einstellungen"].iter_rows(values_only=True))[1:]
            if z[0]} if "Einstellungen" in wb.sheetnames else {}
    cfg = _cfg_aus_tabellen(t, rest, stundentafel)
    kinder = lese_kinder(pfad)
    return cfg, kinder


# ------------------------------------------------------------ Raum- und Lehrpersonenlisten

# Spaltennamen, wie sie in Raumverwaltungen, Hauswart- oder Personallisten üblich sind
LISTEN_ALIASE = {
    "Raeume": {
        "schulhaus": ["schulhaus", "standort", "gebaeude", "schulanlage", "haus", "objekt"],
        "name": ["name", "raum", "raumname", "raumbezeichnung", "bezeichnung", "zimmer", "raumnr", "raumnummer",
                 "zimmernr", "nummer"],
        "typ": ["typ", "raumtyp", "raumart", "art", "nutzung", "kategorie", "verwendung"],
        "plaetze": ["plaetze", "platze", "sitzplaetze", "arbeitsplaetze", "kapazitaet", "kapazitat", "maxpersonen",
                    "personen", "max", "groesse"],
        "anzahl": ["anzahl", "stueck", "stuck", "menge"],
        "nicht_verfuegbar": ["nichtverfuegbar", "nichtverfugbar", "gesperrt", "sperrzeiten", "fremdnutzung",
                             "belegt", "blockiert"],
    },
    "Lehrpersonen": {
        "name": ["name", "lehrperson", "lp", "person", "vollername"],
        "vorname": ["vorname"],
        "nachname": ["nachname", "familienname"],
        "kuerzel": ["kuerzel", "kurzel", "kurzzeichen", "visum", "kz"],
        "faecher": ["faecher", "facher", "fach", "unterrichtsfaecher", "lehrbefaehigung"],
        "jahrgaenge": ["jahrgaenge", "jahrgange", "stufen", "stufe", "klassen"],
        "schulhaeuser": ["schulhaeuser", "schulhauser", "schulhaus", "standort", "einsatzort"],
        "pensum_min": ["pensummin", "minpensum", "minlektionen", "lektionenmin", "mindestpensum"],
        "pensum_max": ["pensummax", "maxpensum", "maxlektionen", "lektionenmax", "hoechstpensum"],
        "pensum": ["pensum", "lektionen", "wochenlektionen"],
        "nicht_verfuegbar": ["nichtverfuegbar", "nichtverfugbar", "freietage", "abwesend", "nichtda",
                             "sperrzeiten", "unverfuegbar"],
        "bisherige_klasse": ["bisherigeklasse", "klassevorjahr", "aktuelleklasse", "klassenlehrperson"],
    },
}

# Raumtypen: gängige Bezeichnungen → Begriffe der Konfiguration
RAUMTYP_ALIASE = {
    "Klassenzimmer": ["klassenzimmer", "schulzimmer", "zimmer", "klassenraum", "unterrichtsraum", "kz", "sz"],
    "Turnhalle": ["turnhalle", "halle", "sporthalle", "mehrzweckhalle", "bewegungundsport", "sport"],
    "TTG-Zimmer": ["ttgzimmer", "ttg", "werkraum", "handarbeit", "textilesgestalten", "technischesgestalten",
                   "werken"],
    "Kindergarten": ["kindergarten", "kg", "kiga", "kindergartenraum"],
    "Gruppenraum": ["gruppenraum", "gruppenzimmer", "foerderraum", "ifraum", "dazraum"],
}

_TAGE = {"montag": "Mo", "mo": "Mo", "dienstag": "Di", "di": "Di", "mittwoch": "Mi", "mi": "Mi",
         "donnerstag": "Do", "do": "Do", "freitag": "Fr", "fr": "Fr"}
_HALBTAGE = {"vormittag": "VM", "vm": "VM", "morgen": "VM", "nachmittag": "NM", "nm": "NM"}


def zeitangabe(text: str) -> str:
    """'Freitag Nachmittag' → 'Fr-NM', 'mo vm 1' → 'Mo-VM-1', 'Mi' → 'Mi'."""
    teile = [t for t in re.split(r"[\s\-_/.,]+", str(text).strip().lower()) if t]
    if not teile or _norm(teile[0]) not in _TAGE:
        raise ValueError(f"Zeitangabe '{text}' nicht verstanden (Beispiele: Mi, Fr-NM, Mo-VM-1, Freitag Nachmittag)")
    aus = [_TAGE[_norm(teile[0])]]
    for t in teile[1:]:
        if _norm(t) in _HALBTAGE and len(aus) == 1:
            aus.append(_HALBTAGE[_norm(t)])
        elif t.isdigit() and len(aus) == 2:
            aus.append(t)
        elif t not in ("ganz", "ganzer", "tag", "lektion", "l"):
            raise ValueError(f"Zeitangabe '{text}' nicht verstanden (Beispiele: Mi, Fr-NM, Mo-VM-1)")
    return "-".join(aus)


def _zeitliste(v: str) -> list[str] | None:
    teile = [x for x in re.split(r"[;\n]+|,(?=\s*[A-Za-zÄÖÜäöü])", str(v or "")) if x.strip()]
    return [zeitangabe(x) for x in teile] or None


def raumtyp(text: str, bekannte: list[str] | tuple[str, ...] = ()) -> str:
    """Raumtyp auf einen bekannten Begriff abbilden; Unbekanntes bleibt wörtlich."""
    n = _norm(text)
    for b in bekannte:
        if _norm(b) == n:
            return b
    for typ, aliase in RAUMTYP_ALIASE.items():
        if n in aliase:
            return next((b for b in bekannte if _norm(b) == _norm(typ)), typ)
    return str(text).strip()


def lese_liste(pfad: str | Path, art: str, spalten: dict[str, str] | None = None,
               raumtypen: list[str] | tuple[str, ...] = ()) -> list[dict]:
    """Raum- oder Lehrpersonenliste (CSV oder Excel, beliebige gängige Spaltennamen) lesen.

    ``art`` ist ``"Raeume"`` oder ``"Lehrpersonen"``. Liefert Zeilen im Format der
    Stammdaten-Tabellen (BLAETTER), bereit für die Konfiguration."""
    if art not in LISTEN_ALIASE:
        raise ValueError(f"Unbekannte Liste '{art}' (möglich: {', '.join(LISTEN_ALIASE)})")
    zeilen = _zeilen(pfad, (art, "Räume" if art == "Raeume" else art))
    if not zeilen:
        return []
    kopf = list(zeilen[0].keys())
    zuordnung: dict[str, str] = {}
    for feld, quelle in (spalten or {}).items():
        if quelle not in kopf:
            raise ValueError(f"Spalte '{quelle}' (für {feld}) fehlt in {pfad}")
        zuordnung[feld] = quelle
    normiert = {_norm(k): k for k in kopf}
    for feld, aliase in LISTEN_ALIASE[art].items():
        if feld not in zuordnung:
            treffer = next((normiert[a] for a in aliase if a in normiert and normiert[a] not in zuordnung.values()),
                           None)
            if treffer:
                zuordnung[feld] = treffer
    pflicht = ["schulhaus", "typ"] if art == "Raeume" else []
    if art == "Lehrpersonen" and "name" not in zuordnung and "nachname" not in zuordnung:
        pflicht.append("name")
    fehlend = [f for f in pflicht if f not in zuordnung]
    if fehlend:
        raise ValueError(f"{Path(pfad).name}: Spalte(n) {', '.join(fehlend)} nicht gefunden "
                         f"(vorhanden: {', '.join(kopf)}). Mit --spalte {fehlend[0]}=<Spaltenname> zuordnen.")
    out = []
    for nr, zeile in enumerate(zeilen, start=2):
        def w(feld: str) -> str:
            return str(zeile.get(zuordnung.get(feld, ""), "") or "").strip()

        try:
            e: dict = {}
            if art == "Raeume":
                if not w("schulhaus") and not w("typ"):
                    continue
                e = {"schulhaus": w("schulhaus"), "typ": raumtyp(w("typ"), raumtypen)}
                if w("name"):
                    e["name"] = w("name")
                for feld in ("anzahl", "plaetze"):
                    if w(feld):
                        e[feld] = _zahl(w(feld).replace("'", "").replace(",", "."))
            else:
                name = w("name") or " ".join(x for x in (w("vorname"), w("nachname")) if x)
                if not name:
                    continue
                e = {"name": name}
                if w("kuerzel"):
                    e["kuerzel"] = w("kuerzel")
                for feld in ("faecher", "jahrgaenge", "schulhaeuser"):
                    if w(feld):
                        e[feld] = [x.strip() for x in re.split(r"[;,/]", w(feld)) if x.strip()]
                if w("pensum") and "-" in w("pensum"):
                    a, b = w("pensum").split("-", 1)
                    e["pensum_min"], e["pensum_max"] = _zahl(a.strip()), _zahl(b.strip())
                elif w("pensum"):
                    e["pensum_min"] = e["pensum_max"] = _zahl(w("pensum"))
                for feld in ("pensum_min", "pensum_max"):
                    if w(feld):
                        e[feld] = _zahl(w(feld))
                if w("bisherige_klasse"):
                    e["bisherige_klasse"] = w("bisherige_klasse")
            if w("nicht_verfuegbar"):
                e["nicht_verfuegbar"] = _zeitliste(w("nicht_verfuegbar"))
        except ValueError as fehler:
            raise ValueError(f"{Path(pfad).name}, Zeile {nr}: {fehler}") from None
        out.append(e)
    return out


def uebernehme_liste(cfg: dict, art: str, zeilen: list[dict], ersetzen: bool = True) -> tuple[dict, list[str]]:
    """Liste in die Konfiguration übernehmen. ``ersetzen``: bisherige Räume der betroffenen
    Schulhäuser bzw. alle Lehrpersonen ersetzen; sonst ergänzen (gleicher Name überschreibt).
    Liefert (neue Konfiguration, Hinweise)."""
    t, extra = _tabellen_aus_cfg(cfg)
    zeilen = json.loads(json.dumps(zeilen))
    hinweise: list[str] = []
    if art == "Raeume":
        haeuser = {sh["name"] for sh in t["Schulhaeuser"]}
        unbekannt = sorted({z["schulhaus"] for z in zeilen} - haeuser)
        if unbekannt:
            raise ValueError(f"Unbekannte Schulhäuser: {', '.join(unbekannt)} – zuerst unter Schulhäuser erfassen "
                             f"(vorhanden: {', '.join(sorted(haeuser))}).")
        betroffen = {z["schulhaus"] for z in zeilen}
        alt = [r for r in t["Raeume"] if not (ersetzen and r["schulhaus"] in betroffen)]
        schluessel = {(z["schulhaus"], z.get("name")) for z in zeilen if z.get("name")}
        t["Raeume"] = [r for r in alt if (r["schulhaus"], r.get("name")) not in schluessel] + zeilen
        benoetigt = {v.get("raumtyp", "Klassenzimmer") for v in t["Klassentypen"]} | \
                    {f.get("raumtyp") for f in t["Faecher"] if f.get("raumtyp")}
        neu = sorted({z["typ"] for z in zeilen} - benoetigt)
        if neu:
            hinweise.append(f"Raumtyp(en) {', '.join(neu)} werden von keinem Fach und keinem Klassentyp verlangt "
                            "– sie werden nicht verplant.")
        plaetze = sum(int(z.get("anzahl", 1)) for z in zeilen)
        hinweise.insert(0, f"{plaetze} Räume in {len(betroffen)} {'Schulhaus' if len(betroffen) == 1 else 'Schulhäusern'} "
                           "übernommen.")
    elif art == "Lehrpersonen":
        if ersetzen:
            t["Lehrpersonen"] = []
        namen = {z["name"] for z in zeilen}
        t["Lehrpersonen"] = [lp for lp in t["Lehrpersonen"] if lp["name"] not in namen] + zeilen
        faecher = {f["name"] for f in t["Faecher"]} | {"IF", "DaZ Anfang", "DaZ Aufbau", "IS"}
        fremd = sorted({f for z in zeilen for f in z.get("faecher", [])} - faecher)
        if fremd:
            hinweise.append(f"Unbekannte Fächer: {', '.join(fremd)} – Schreibweise wie unter Fächer verwenden.")
        ohne = [z["name"] for z in zeilen if not z.get("faecher")]
        if ohne:
            hinweise.append(f"{len(ohne)} Lehrpersonen ohne Fächer (werden nicht eingesetzt): {', '.join(ohne[:5])}"
                            f"{' …' if len(ohne) > 5 else ''}")
        hinweise.insert(0, f"{len(zeilen)} Lehrpersonen übernommen.")
    else:
        raise ValueError(f"Unbekannte Liste '{art}'")
    return _cfg_aus_tabellen(t, extra["rest"], extra["stundentafel"]), hinweise


def raumtypen(cfg: dict) -> list[str]:
    """Alle in der Konfiguration verwendeten Raumtypen (für die Zuordnung beim Import)."""
    t, _ = _tabellen_aus_cfg(cfg)
    return sorted({r["typ"] for r in t["Raeume"]} | {v.get("raumtyp", "Klassenzimmer") for v in t["Klassentypen"]}
                  | {f["raumtyp"] for f in t["Faecher"] if f.get("raumtyp")})



def schreibe_vorlage(ordner: str | Path) -> list[Path]:
    """Leere, ausgefüllte Beispiel-Vorlagen zum Erfassen eigener Daten."""
    ordner = Path(ordner)
    ordner.mkdir(parents=True, exist_ok=True)
    from .daten import beispiel_ordner
    beispiel = beispiel_ordner() / "vorlage"
    dateien = []
    for name in ("schule.yaml", "kinder.csv", "raumliste.csv"):
        ziel = ordner / name
        ziel.write_text((beispiel / name).read_text(encoding="utf-8"), encoding="utf-8")
        dateien.append(ziel)
    cfg = lese_rohkonfiguration(ordner / "schule.yaml")
    exportiere_stammdaten(cfg, lese_kinder(ordner / "kinder.csv"), ordner / "stammdaten.xlsx")
    dateien.append(ordner / "stammdaten.xlsx")
    return dateien
